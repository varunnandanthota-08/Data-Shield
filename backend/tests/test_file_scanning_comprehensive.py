import io
import json
import pytest
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from app.main import app
from app.core.config import settings
from app.engine.file_extractor import (
    extract_file_content,
    extract_pdf_text,
    extract_docx_text,
    FileExtractionError,
    MAX_FILE_SIZE_BYTES,
    MAX_EXTRACTED_CHARS,
)
from app.db.repository import ScanRecordRepository
from app.ai.service import GeminiContextAnalysisService

client = TestClient(app)

SYNTHETIC_AWS_KEY = "AKIAIOSFODNN7EXAMPLE"
SYNTHETIC_SSN = "123-45-6789"
SYNTHETIC_CARD = "4111111111111111"

# Helper to generate a minimal, valid in-memory PDF with readable text
def create_minimal_pdf_bytes(text_line: str) -> bytes:
    # Clean text line for PDF stream
    clean_line = text_line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream_content = f"BT\n/F1 12 Tf\n50 700 Td\n({clean_line}) Tj\nET\n".encode("latin-1")
    stream_len = len(stream_content)

    pdf = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
        b"4 0 obj\n<< /Length " + str(stream_len).encode("ascii") + b" >>\nstream\n"
        + stream_content +
        b"endstream\nendobj\n"
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000224 00000 n \n"
        b"0000000348 00000 n \n"
        b"trailer\n<< /Size 6 /Root 1 0 R >>\n"
        b"startxref\n450\n%%EOF\n"
    )
    return pdf


# Helper to generate a valid in-memory DOCX with paragraphs and tables
def create_minimal_docx_bytes(paragraph_text: str, table_rows: list[list[str]]) -> bytes:
    import docx
    doc = docx.Document()
    doc.add_heading("Confidential Intake Document", level=1)
    doc.add_paragraph(paragraph_text)
    
    if table_rows:
        num_rows = len(table_rows)
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=num_rows, cols=num_cols)
        for r_idx, row_data in enumerate(table_rows):
            for c_idx, cell_data in enumerate(row_data):
                table.rows[r_idx].cells[c_idx].text = cell_data

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestFileExtractionCore:
    """Step 2 & 3: File content validation, safe in-memory extraction, and security checks."""

    def test_extract_plain_text_success(self):
        content = b"system_metric=cpu_load\nthreshold=0.85\nstatus=nominal"
        res = extract_file_content("metrics.conf", content)
        assert res.filename == "metrics.conf"
        assert res.file_type == "configuration"
        assert "threshold=0.85" in res.text_content
        assert res.size_bytes == len(content)
        assert res.is_truncated is False

    def test_extract_source_code_success(self):
        content = b"export function auth() { return true; }"
        res = extract_file_content("auth.ts", content)
        assert res.filename == "auth.ts"
        assert res.file_type == "source_code"
        assert "export function auth" in res.text_content

    def test_extract_pdf_document_success(self):
        pdf_bytes = create_minimal_pdf_bytes(f"Report: AWS_KEY={SYNTHETIC_AWS_KEY}")
        res = extract_file_content("audit_findings.pdf", pdf_bytes)
        assert res.filename == "audit_findings.pdf"
        assert res.file_type == "document"
        assert SYNTHETIC_AWS_KEY in res.text_content
        assert res.extraction_note is not None
        assert "PDF pages" in res.extraction_note

    def test_extract_docx_document_success(self):
        docx_bytes = create_minimal_docx_bytes(
            paragraph_text=f"Employee SSN: {SYNTHETIC_SSN}",
            table_rows=[["Account", "Card"], ["Jane", SYNTHETIC_CARD]],
        )
        res = extract_file_content("payroll.docx", docx_bytes)
        assert res.filename == "payroll.docx"
        assert res.file_type == "document"
        assert SYNTHETIC_SSN in res.text_content
        assert SYNTHETIC_CARD in res.text_content
        assert res.extraction_note is not None
        assert "paragraph and table text" in res.extraction_note

    def test_reject_spoofed_pdf_without_magic_header(self):
        fake_pdf = b"This is just a plain text file pretending to be a PDF."
        with pytest.raises(FileExtractionError) as exc_info:
            extract_file_content("spoofed.pdf", fake_pdf)
        assert exc_info.value.status_code == 400
        assert "valid PDF signature" in exc_info.value.message

    def test_reject_spoofed_docx_without_zip_header(self):
        fake_docx = b"Plain text file renamed to .docx."
        with pytest.raises(FileExtractionError) as exc_info:
            extract_file_content("spoofed.docx", fake_docx)
        assert exc_info.value.status_code == 400
        assert "valid DOCX container header" in exc_info.value.message

    def test_reject_null_bytes_in_text_file(self):
        binary_payload = b"header\x00\x00binary_data\xff"
        with pytest.raises(FileExtractionError) as exc_info:
            extract_file_content("corrupt.txt", binary_payload)
        assert exc_info.value.status_code == 400
        assert "binary null bytes" in exc_info.value.message

    def test_reject_empty_file_zero_bytes(self):
        with pytest.raises(FileExtractionError) as exc_info:
            extract_file_content("empty.log", b"")
        assert exc_info.value.status_code == 400
        assert "empty (0 bytes)" in exc_info.value.message

    def test_reject_dangerous_binary_extensions(self):
        for bad_ext in [".exe", ".dll", ".sh", ".bat", ".zip", ".tar", ".jar"]:
            with pytest.raises(FileExtractionError) as exc_info:
                extract_file_content(f"payload{bad_ext}", b"data")
            assert exc_info.value.status_code == 400
            assert "rejected for security reasons" in exc_info.value.message

    def test_reject_unsupported_extensions(self):
        for unsupp in [".xyz", ".mp3", ".avi", ".dat"]:
            with pytest.raises(FileExtractionError) as exc_info:
                extract_file_content(f"file{unsupp}", b"some content")
            assert exc_info.value.status_code == 415
            assert "Unsupported file format" in exc_info.value.message

    def test_reject_file_exceeding_size_limit(self):
        oversized = b"A" * (MAX_FILE_SIZE_BYTES + 1024)
        with pytest.raises(FileExtractionError) as exc_info:
            extract_file_content("large.log", oversized)
        assert exc_info.value.status_code == 413
        assert "exceeds maximum permitted limit" in exc_info.value.message

    def test_content_truncation_when_exceeding_max_chars(self):
        # Create a text file with more than MAX_EXTRACTED_CHARS
        huge_text = ("line of log text with no secrets.\n" * 25000).encode("utf-8")
        res = extract_file_content("huge_log.txt", huge_text)
        assert len(res.text_content) <= MAX_EXTRACTED_CHARS
        assert res.is_truncated is True
        assert res.extraction_note is not None
        assert "truncated" in res.extraction_note


class TestFileScanningAPIWorkflows:
    """End-to-end API scan testing: DLP inspection, hard-block enforcement, redaction, audit, and AI fallback."""

    @pytest.fixture(autouse=True)
    def mock_gemini_fast_response(self):
        """Mock Gemini client during test runs to ensure instant, deterministic execution."""
        mock_resp = MagicMock()
        mock_resp.text = json.dumps({
            "incident_summary": "Document scanning completed safely.",
            "relevant_risk_factors": ["Sensitive credentials detected in uploaded document."],
            "policy_rationales": {
                "POL-001": "Deterministic blocking policy enforced for credentials.",
            },
            "ambiguities_and_uncertainty": [],
            "remediation_recommendations": [
                {
                    "id": "REM-01",
                    "action_type": "ROTATE_CREDENTIAL",
                    "title": "Rotate Uploaded Credential",
                    "explanation": "Credentials found in uploaded files should be revoked immediately.",
                    "reason": "Exposed in document payload",
                    "expected_security_benefit": "Neutralizes exposed access token",
                    "requires_human_approval": False,
                }
            ],
            "human_review_recommended": False,
        })
        with patch.object(GeminiContextAnalysisService, "_get_client") as mock_get_client:
            mock_client = MagicMock()
            mock_client.models.generate_content.return_value = mock_resp
            mock_get_client.return_value = mock_client
            yield mock_client

    def test_scan_pdf_with_secret_triggers_mandatory_hard_block(self):
        pdf_bytes = create_minimal_pdf_bytes(f"Deploy config: AWS_SECRET={SYNTHETIC_AWS_KEY}")
        files = {"file": ("production_keys.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        data = {
            "context_json": json.dumps({
                "destination": "PUBLIC_LLM",
                "recipient": "PUBLIC",
                "classification": "CONFIDENTIAL",
                "actor_id": "compliance_tester",
                "org_id": "default",
            })
        }

        response = client.post("/api/v1/scan/file", files=files, data=data)
        assert response.status_code == 200
        res = response.json()

        assert res["source_type"] == "file"
        assert res["filename"] == "production_keys.pdf"
        assert res["policy_decision"] == "BLOCK"
        assert res["risk_band"] == "CRITICAL"
        assert res["counts"]["total"] >= 1
        # Hard invariant: raw secret never appears in redacted text
        assert SYNTHETIC_AWS_KEY not in res["redacted_text"]
        assert "[REDACTED_AWS_KEY]" in res["redacted_text"]
        # Extraction note present
        assert "PDF pages" in (res.get("extraction_note") or "")
        # Audit proof
        assert res.get("audit_event_id") is not None
        assert res.get("audit_event_hash") is not None

    def test_scan_docx_with_pii_and_card_redaction(self):
        docx_bytes = create_minimal_docx_bytes(
            paragraph_text=f"Customer Intake: SSN={SYNTHETIC_SSN}",
            table_rows=[["Holder", "Number"], ["Alice", SYNTHETIC_CARD]],
        )
        files = {"file": ("intake.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {
            "context_json": json.dumps({
                "destination": "EXTERNAL_SAAS",
                "recipient": "THIRD_PARTY_CONTRACTOR",
                "classification": "CONFIDENTIAL",
                "actor_id": "ops_analyst",
                "org_id": "default",
            })
        }

        response = client.post("/api/v1/scan/file", files=files, data=data)
        assert response.status_code == 200
        res = response.json()

        assert res["source_type"] == "file"
        assert res["filename"] == "intake.docx"
        assert res["counts"]["total"] >= 2
        # Zero raw sensitive data in redacted document output
        assert SYNTHETIC_SSN not in res["redacted_text"]
        assert SYNTHETIC_CARD not in res["redacted_text"]
        assert "[REDACTED_SSN]" in res["redacted_text"]
        assert "[REDACTED_CARD]" in res["redacted_text"]
        assert "paragraph and table text" in (res.get("extraction_note") or "")

    def test_scan_file_oversized_returns_413(self):
        huge_bytes = b"0" * (MAX_FILE_SIZE_BYTES + 500)
        files = {"file": ("big_log.txt", io.BytesIO(huge_bytes), "text/plain")}
        response = client.post("/api/v1/scan/file", files=files)
        assert response.status_code == 413
        assert "exceeds maximum permitted size" in response.json()["detail"]

    def test_scan_file_unsupported_format_returns_415(self):
        files = {"file": ("unsupported.dat", io.BytesIO(b"some data"), "application/octet-stream")}
        response = client.post("/api/v1/scan/file", files=files)
        assert response.status_code == 415
        assert "Unsupported file format" in response.json()["detail"]

    def test_scan_file_corrupt_pdf_returns_400(self):
        corrupt_pdf = b"Not a real PDF structure"
        files = {"file": ("corrupt.pdf", io.BytesIO(corrupt_pdf), "application/pdf")}
        response = client.post("/api/v1/scan/file", files=files)
        assert response.status_code == 400
        assert "valid PDF signature" in response.json()["detail"]

    def test_scan_file_empty_returns_400(self):
        files = {"file": ("empty.json", io.BytesIO(b""), "application/json")}
        response = client.post("/api/v1/scan/file", files=files)
        assert response.status_code == 400
        assert "empty (0 bytes)" in response.json()["detail"]

    def test_scan_file_zero_retention_in_audit_and_scan_records(self):
        pdf_bytes = create_minimal_pdf_bytes(f"Confidential AWS Key: {SYNTHETIC_AWS_KEY}")
        files = {"file": ("audit_leak.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        data = {
            "context_json": json.dumps({
                "destination": "PUBLIC_LLM",
                "recipient": "PUBLIC",
                "classification": "CONFIDENTIAL",
                "actor_id": "auditor",
                "org_id": "default",
            })
        }

        response = client.post("/api/v1/scan/file", files=files, data=data)
        assert response.status_code == 200
        res = response.json()
        scan_id = res["scan_id"]

        # Check repository scan record - must never contain raw secret
        repo = ScanRecordRepository()
        rec = repo.get_scan_record(scan_id)
        assert rec is not None
        assert SYNTHETIC_AWS_KEY not in json.dumps(rec)

        # Check audit event - must never contain raw secret
        audit_event_id = res["audit_event_id"]
        audit_res = client.get("/api/v1/audit/events?limit=20&offset=0")
        assert audit_res.status_code == 200
        matched = next((e for e in audit_res.json()["items"] if e["event_id"] == audit_event_id), None)
        assert matched is not None
        assert matched["event_type"] == "FILE_SCAN_EVALUATION"
        assert SYNTHETIC_AWS_KEY not in json.dumps(matched)
        assert matched["sanitized_metadata"]["source"]["filename"] == "audit_leak.pdf"
        assert matched["sanitized_metadata"]["source"]["source_type"] == "file"

    def test_gemini_unavailable_fallback_deterministic_scan_preserves_policy(self):
        """When Gemini is unconfigured or unavailable, deterministic policy evaluation and audit hashing continue uninterrupted."""
        with patch.object(GeminiContextAnalysisService, "is_configured", new_callable=lambda: property(lambda self: False)):
            content = f"Database password: {SYNTHETIC_AWS_KEY}".encode("utf-8")
            files = {"file": ("db.env", io.BytesIO(content), "text/plain")}
            data = {
                "context_json": json.dumps({
                    "destination": "PUBLIC_LLM",
                    "recipient": "PUBLIC",
                    "classification": "CONFIDENTIAL",
                    "actor_id": "analyst",
                    "org_id": "default",
                })
            }

            response = client.post("/api/v1/scan/file", files=files, data=data)
            assert response.status_code == 200
            res = response.json()
            assert res["policy_decision"] == "BLOCK"
            assert res["ai_analysis"]["ai_status"] == "UNAVAILABLE"
            assert res["ai_analysis"]["is_ai_assisted"] is False
            assert "deterministic-fallback" in res["ai_analysis"]["model_identifier"]

    def test_gemini_invocation_error_fallback_preserves_policy(self):
        """When Gemini invocation fails due to network or provider error, deterministic fallback is safely engaged."""
        mock_client = MagicMock()
        mock_client.models.generate_content.side_effect = Exception("503 Service Unavailable")
        with patch.object(GeminiContextAnalysisService, "_get_client", return_value=mock_client):
            content = f"AWS_KEY={SYNTHETIC_AWS_KEY}".encode("utf-8")
            files = {"file": ("config.env", io.BytesIO(content), "text/plain")}
            data = {
                "context_json": json.dumps({
                    "destination": "PUBLIC_LLM",
                    "recipient": "PUBLIC",
                    "classification": "CONFIDENTIAL",
                    "actor_id": "analyst",
                    "org_id": "default",
                })
            }

            response = client.post("/api/v1/scan/file", files=files, data=data)
            assert response.status_code == 200
            res = response.json()
            assert res["policy_decision"] == "BLOCK"
            assert res["ai_analysis"]["ai_status"] == "FALLBACK"
            assert res["ai_analysis"]["is_ai_assisted"] is False
            assert "deterministic-fallback" in res["ai_analysis"]["model_identifier"]
