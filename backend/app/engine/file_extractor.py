import io
import os
import re
from typing import Dict, List, Optional, Set, Tuple
from app.core.config import settings

# Maximum permitted uploaded file size: default 5 MB (configurable via MAX_UPLOAD_SIZE_BYTES)
MAX_FILE_SIZE_BYTES = settings.MAX_UPLOAD_SIZE_BYTES

# Maximum permitted extracted characters to scan: default 500,000 chars (configurable via MAX_EXTRACTED_CHARS)
MAX_EXTRACTED_CHARS = settings.MAX_EXTRACTED_CHARS

# Supported extension sets
TEXT_AND_DATA_EXTENSIONS: Set[str] = {
    ".txt", ".log", ".csv", ".json", ".yaml", ".yml", ".xml",
}

CODE_AND_MARKUP_EXTENSIONS: Set[str] = {
    ".md", ".html", ".css", ".js", ".ts", ".py", ".java",
}

CONFIG_AND_SYSTEM_EXTENSIONS: Set[str] = {
    ".sql", ".env", ".ini", ".conf", ".properties",
}

DOCUMENT_EXTENSIONS: Set[str] = {
    ".pdf", ".docx",
}

ALL_SUPPORTED_EXTENSIONS: Set[str] = (
    TEXT_AND_DATA_EXTENSIONS
    | CODE_AND_MARKUP_EXTENSIONS
    | CONFIG_AND_SYSTEM_EXTENSIONS
    | DOCUMENT_EXTENSIONS
)

# Known dangerous executable extensions
DANGEROUS_BINARY_EXTENSIONS: Set[str] = {
    ".exe", ".dll", ".so", ".dylib", ".bin", ".elf", ".sh", ".bat", ".cmd", ".vbs",
    ".ps1", ".jar", ".war", ".class", ".zip", ".tar", ".gz", ".7z", ".rar", ".iso",
}


class FileExtractionError(Exception):
    """Safe, non-leaking exception raised when file extraction or validation fails."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


class ExtractedFilePayload:
    """Represents safely extracted and bounded document content."""
    def __init__(
        self,
        filename: str,
        extension: str,
        size_bytes: int,
        text_content: str,
        file_type: str = "plain_text",
        is_truncated: bool = False,
        extraction_note: Optional[str] = None,
    ):
        self.filename = filename
        self.extension = extension
        self.size_bytes = size_bytes
        self.text_content = text_content
        self.file_type = file_type
        self.is_truncated = is_truncated
        self.extraction_note = extraction_note


def sanitize_filename(filename: Optional[str]) -> str:
    """Sanitizes user-provided filename, stripping directory traversal components."""
    if not filename:
        return "unnamed_upload"
    # Take basename only
    base = os.path.basename(filename.strip().replace("\\", "/"))
    # Remove control characters
    base = re.sub(r"[\x00-\x1f\x7f]", "", base)
    return base if base else "unnamed_upload"


def validate_file_metadata(filename: str, size_bytes: int) -> str:
    """Validates filename extension and size against strict allowlists."""
    if size_bytes > MAX_FILE_SIZE_BYTES:
        raise FileExtractionError(
            f"File size ({size_bytes:,} bytes) exceeds maximum permitted limit of {MAX_FILE_SIZE_BYTES:,} bytes (5 MB).",
            status_code=413,
        )

    _, ext = os.path.splitext(filename.lower())
    if not ext:
        raise FileExtractionError(
            "File extension missing. Please upload a file with an explicit supported extension (e.g., .txt, .json, .log, .pdf).",
            status_code=400,
        )

    if ext in DANGEROUS_BINARY_EXTENSIONS:
        raise FileExtractionError(
            f"File extension '{ext}' is an executable or archive format and is rejected for security reasons.",
            status_code=400,
        )

    if ext not in ALL_SUPPORTED_EXTENSIONS:
        supported_list = ", ".join(sorted(ALL_SUPPORTED_EXTENSIONS))
        raise FileExtractionError(
            f"Unsupported file format '{ext}'. Supported formats: {supported_list}.",
            status_code=415,
        )

    return ext


def extract_plain_text(raw_bytes: bytes, filename: str) -> str:
    """Decodes plain text, source code, and configuration files with safe encodings."""
    # Check for binary null bytes which indicate non-text files
    if b"\x00" in raw_bytes[:4096]:
        raise FileExtractionError(
            f"File '{filename}' contains binary null bytes and cannot be scanned as plain text.",
            status_code=400,
        )

    # Try UTF-8 first
    try:
        return raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        pass

    # Fallback to latin-1
    try:
        return raw_bytes.decode("latin-1")
    except Exception as e:
        raise FileExtractionError(
            f"Failed to decode text content from '{filename}'. Please ensure the file uses standard text encoding (UTF-8).",
            status_code=400,
        )


def extract_pdf_text(raw_bytes: bytes, filename: str) -> Tuple[str, Optional[str]]:
    """Extracts text from PDF documents using pypdf."""
    # Content-based magic bytes validation: PDF must contain %PDF- in header
    if b"%PDF-" not in raw_bytes[:1024]:
        raise FileExtractionError(
            f"File '{filename}' does not contain a valid PDF signature.",
            status_code=400,
        )

    try:
        import pypdf
    except ImportError:
        raise FileExtractionError(
            "PDF text extraction library is not available on this server.",
            status_code=500,
        )

    try:
        pdf_stream = io.BytesIO(raw_bytes)
        reader = pypdf.PdfReader(pdf_stream)

        if reader.is_encrypted:
            try:
                # Try decrypting with empty password
                reader.decrypt("")
            except Exception:
                raise FileExtractionError(
                    f"PDF document '{filename}' is password-protected or encrypted and cannot be parsed.",
                    status_code=400,
                )

        extracted_pages: List[str] = []
        total_pages = len(reader.pages)
        max_pages = min(total_pages, 100)

        for page_idx in range(max_pages):
            page = reader.pages[page_idx]
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text)

        full_text = "\n\n".join(extracted_pages)
        note = (
            f"Extracted text from {max_pages}/{total_pages} PDF pages. "
            "Note: Scanned images and embedded non-text objects require OCR and are not extracted."
        )

        if not full_text.strip():
            raise FileExtractionError(
                f"No readable text could be extracted from PDF '{filename}'. "
                "The file may contain only scanned images or vector graphics.",
                status_code=400,
            )

        return full_text, note

    except FileExtractionError:
        raise
    except Exception as e:
        raise FileExtractionError(
            f"Failed to parse PDF document '{filename}'. File may be corrupt or formatted improperly.",
            status_code=400,
        )


def extract_docx_text(raw_bytes: bytes, filename: str) -> Tuple[str, Optional[str]]:
    """Extracts text from DOCX documents using python-docx."""
    # Content-based magic bytes validation: DOCX must start with PK zip container header
    if not (raw_bytes.startswith(b"PK\x03\x04") or raw_bytes.startswith(b"PK\x05\x06")):
        raise FileExtractionError(
            f"File '{filename}' does not contain a valid DOCX container header.",
            status_code=400,
        )

    try:
        import docx
    except ImportError:
        raise FileExtractionError(
            "DOCX text extraction library is not available on this server.",
            status_code=500,
        )

    try:
        docx_stream = io.BytesIO(raw_bytes)
        doc = docx.Document(docx_stream)

        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)

        all_text = "\n".join(paragraphs_text + table_text)
        note = "Extracted paragraph and table text from DOCX. Embedded images and comments are excluded."

        if not all_text.strip():
            raise FileExtractionError(
                f"No readable text found in DOCX file '{filename}'. Document appears empty.",
                status_code=400,
            )

        return all_text, note

    except FileExtractionError:
        raise
    except Exception as e:
        raise FileExtractionError(
            f"Failed to parse DOCX document '{filename}'. File may be corrupt or invalid.",
            status_code=400,
        )


def extract_file_content(filename: str, raw_bytes: bytes) -> ExtractedFilePayload:
    """Validates and extracts plain text from supported uploaded files.
    
    Guarantees:
    - Rejects empty files.
    - Rejects files exceeding MAX_FILE_SIZE_BYTES.
    - Decodes only supported file formats.
    - Truncates extracted text to MAX_EXTRACTED_CHARS.
    - Sanitizes filenames and returns structured ExtractedFilePayload.
    """
    safe_name = sanitize_filename(filename)
    size_bytes = len(raw_bytes)

    if size_bytes == 0:
        raise FileExtractionError(
            f"Uploaded file '{safe_name}' is empty (0 bytes). Please upload a valid document.",
            status_code=400,
        )

    ext = validate_file_metadata(safe_name, size_bytes)
    extraction_note: Optional[str] = None

    if ext in CODE_AND_MARKUP_EXTENSIONS:
        file_type = "source_code"
    elif ext in CONFIG_AND_SYSTEM_EXTENSIONS:
        file_type = "configuration"
    elif ext in DOCUMENT_EXTENSIONS:
        file_type = "document"
    else:
        file_type = "plain_text"

    if ext in (TEXT_AND_DATA_EXTENSIONS | CODE_AND_MARKUP_EXTENSIONS | CONFIG_AND_SYSTEM_EXTENSIONS):
        raw_text = extract_plain_text(raw_bytes, safe_name)
    elif ext == ".pdf":
        raw_text, extraction_note = extract_pdf_text(raw_bytes, safe_name)
    elif ext == ".docx":
        raw_text, extraction_note = extract_docx_text(raw_bytes, safe_name)
    else:
        raise FileExtractionError(
            f"Unsupported file format '{ext}'.",
            status_code=415,
        )

    if not raw_text.strip():
        raise FileExtractionError(
            f"File '{safe_name}' contains no readable text content to scan.",
            status_code=400,
        )

    is_truncated = False
    if len(raw_text) > MAX_EXTRACTED_CHARS:
        raw_text = raw_text[:MAX_EXTRACTED_CHARS]
        is_truncated = True
        trunc_note = f"Content truncated to {MAX_EXTRACTED_CHARS:,} characters for bounded security scanning."
        extraction_note = f"{extraction_note} | {trunc_note}" if extraction_note else trunc_note

    return ExtractedFilePayload(
        filename=safe_name,
        extension=ext,
        size_bytes=size_bytes,
        text_content=raw_text,
        file_type=file_type,
        is_truncated=is_truncated,
        extraction_note=extraction_note,
    )
